"""
Comprehensive Regression & Automated Ingestion Integrity Test Suite for CareerLens AI
Tests all 16 integrity, section preservation, and content loss scenarios.
"""

import io
import os
import sys
import unittest
import docx
import pymupdf
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from database import init_db
from services.resume_parser import (
    clean_text,
    extract_sections,
    extract_sections_detailed,
    extract_text_from_file,
    extract_text_from_docx,
    extract_text_from_pdf,
    ResumeParserError
)
from services.ai_service import ai_service
from services.integrity_service import integrity_service
from services.matching_service import matching_service
from main import app


def build_test_cv_docx_bytes() -> bytes:
    """Builds synthetic CareerLens_Test_CV.docx with all expected sections and projects."""
    doc = docx.Document()
    doc.add_heading("Candidate Name", level=0)
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Full stack software engineer with 3+ years of experience designing and deploying scalable REST APIs and data processing systems.")

    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, FastAPI, React, TypeScript, PostgreSQL, SQL, Docker, Postman, pytest")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Science, Tech University, 2020 - 2024")

    doc.add_heading("Additional Information", level=1)
    doc.add_paragraph("Certifications: AWS Certified Solutions Architect, Certified Scrum Master")

    doc.add_heading("Projects", level=1)
    
    # Project 1: TaskFlow
    doc.add_paragraph("TaskFlow — Task Management Web Application")
    doc.add_paragraph("React and TypeScript frontend", style="List Bullet")
    doc.add_paragraph("Python and FastAPI REST endpoints", style="List Bullet")
    doc.add_paragraph("PostgreSQL storage", style="List Bullet")
    doc.add_paragraph("Postman and pytest testing", style="List Bullet")

    # Project 2: Sales Insights Dashboard
    doc.add_paragraph("Sales Insights Dashboard")
    doc.add_paragraph("Python data-cleaning scripts", style="List Bullet")
    doc.add_paragraph("SQL aggregation queries", style="List Bullet")
    doc.add_paragraph("Sales metrics and dashboard visualizations", style="List Bullet")

    # Project 3: Support Ticket API
    doc.add_paragraph("Support Ticket API")
    doc.add_paragraph("REST API development", style="List Bullet")
    doc.add_paragraph("Ticket status and priority fields", style="List Bullet")
    doc.add_paragraph("Request validation and consistent error responses", style="List Bullet")
    doc.add_paragraph("Basic role-based access rules and API testing", style="List Bullet")

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TestResumeIngestionIntegrity(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        init_db()
        cls.client = TestClient(app)
        from services.auth_service import get_test_token
        cls.client.headers["Authorization"] = f"Bearer {get_test_token('demo_user', 'candidate@careerlens.ai')}"

    # 1. DOCX containing all three expected projects
    def test_01_docx_all_three_expected_projects(self):
        docx_bytes = build_test_cv_docx_bytes()
        text, ftype = extract_text_from_file("CareerLens_Test_CV.docx", docx_bytes)
        self.assertEqual(ftype, "docx")

        analysis = ai_service.analyze_resume(text, filename="CareerLens_Test_CV.docx", file_type="docx")
        profile = analysis["extracted_profile"]
        projects = profile["projects"]

        # Assert all 3 projects are retained
        self.assertEqual(len(projects), 3, f"Expected 3 projects, found {len(projects)}")
        
        titles = [p["title"] for p in projects]
        self.assertTrue(any("TaskFlow" in t for t in titles), f"TaskFlow not found in {titles}")
        self.assertTrue(any("Sales Insights" in t for t in titles), f"Sales Insights not found in {titles}")
        self.assertTrue(any("Support Ticket" in t for t in titles), f"Support Ticket API not found in {titles}")

        # Assert technologies remain associated with the correct projects
        taskflow = next(p for p in projects if "TaskFlow" in p["title"])
        sales = next(p for p in projects if "Sales Insights" in p["title"])
        ticket = next(p for p in projects if "Support Ticket" in p["title"])

        self.assertTrue(any(t in taskflow["technologies"] for t in ["React", "TypeScript", "FastAPI", "PostgreSQL"]))
        self.assertTrue(any(t in sales["technologies"] for t in ["Python", "SQL"]))
        self.assertTrue(any("REST" in t for t in ticket["technologies"]))

        # Assert bullets remain associated with the correct project
        self.assertTrue(any("pytest" in ev.lower() for ev in taskflow["evidence"]))
        self.assertTrue(any("aggregation" in ev.lower() for ev in sales["evidence"]))
        self.assertTrue(any("role-based" in ev.lower() for ev in ticket["evidence"]))

        # Assert Additional Information is preserved
        unrec = profile.get("unrecognized_sections", {})
        self.assertTrue("Additional Information" in unrec or any("Additional" in k for k in unrec))

    # 2. DOCX with multiple projects and multiple bullets per project
    def test_02_docx_multiple_projects_and_bullets(self):
        docx_bytes = build_test_cv_docx_bytes()
        text, _ = extract_text_from_file("test.docx", docx_bytes)
        analysis = ai_service.analyze_resume(text)
        projects = analysis["extracted_profile"]["projects"]

        for p in projects:
            self.assertGreaterEqual(len(p["evidence"]), 3, f"Project {p['title']} has fewer than 3 bullets")
            for ev in p["evidence"]:
                self.assertFalse(ev.startswith("•"), "Bullet prefixes should be normalized from evidence text")

    # 3. DOCX with unusual or alternate section headings
    def test_03_docx_unusual_or_alternate_headings(self):
        doc = docx.Document()
        doc.add_heading("Candidate Person", level=0)
        doc.add_heading("Professional Profile", level=1)
        doc.add_paragraph("Skilled developer")
        doc.add_heading("Core Competencies", level=1)
        doc.add_paragraph("Python, Docker, Kubernetes")
        doc.add_heading("Career History", level=1)
        doc.add_paragraph("Software Engineer at TechCorp\n• Built microservices using Python")
        doc.add_heading("Academic Projects", level=1)
        doc.add_paragraph("Distributed Key-Value Store\n• Implemented Raft consensus in Python")
        doc.add_heading("Educational Qualifications", level=1)
        doc.add_paragraph("B.Tech in Information Technology, 2021")

        bio = io.BytesIO()
        doc.save(bio)
        text, _ = extract_text_from_file("alt_headings.docx", bio.getvalue())
        sections = extract_sections(text)

        self.assertTrue(len(sections["summary"]) > 0, "Summary should match 'Professional Profile'")
        self.assertTrue(len(sections["skills"]) > 0, "Skills should match 'Core Competencies'")
        self.assertTrue(len(sections["experience"]) > 0, "Experience should match 'Career History'")
        self.assertTrue(len(sections["projects"]) > 0, "Projects should match 'Academic Projects'")
        self.assertTrue(len(sections["education"]) > 0, "Education should match 'Educational Qualifications'")

    # 4. DOCX with sections in a different order
    def test_04_docx_sections_different_order(self):
        doc = docx.Document()
        doc.add_heading("Jane Doe", level=0)
        # Order: Education -> Projects -> Skills -> Summary
        doc.add_heading("Education", level=1)
        doc.add_paragraph("M.S. in Computer Science, Stanford, 2022")
        doc.add_heading("Projects", level=1)
        doc.add_paragraph("Cloud Pipeline\n• Built real-time stream processing")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Go, Python, AWS")
        doc.add_heading("Summary", level=1)
        doc.add_paragraph("Backend systems specialist.")

        bio = io.BytesIO()
        doc.save(bio)
        text, _ = extract_text_from_file("order.docx", bio.getvalue())
        analysis = ai_service.analyze_resume(text)
        profile = analysis["extracted_profile"]

        self.assertEqual(len(profile["education"]), 1)
        self.assertEqual(len(profile["projects"]), 1)
        self.assertIn("Python", profile["skills"])
        self.assertEqual(analysis["extraction_status"], "success")

    # 5. Resume with no Projects section
    def test_05_resume_with_no_projects_section(self):
        text = """John Consultant
        Professional Summary:
        Experienced IT management consultant.
        
        Work Experience:
        Senior IT Analyst at Deloitte
        • Managed client enterprise ERP migrations.
        
        Education:
        B.A. in Economics, 2018
        
        Skills:
        Agile, Scrum, Stakeholder Management
        """
        analysis = ai_service.analyze_resume(text, "No_Projects.txt", file_type="txt")
        self.assertEqual(len(analysis["extracted_profile"]["projects"]), 0)
        self.assertIn(analysis["extraction_status"], ["success", "warning"])
        # Should NOT fail or complain of content loss since source has no projects
        self.assertFalse(any("Project section detected" in w for w in analysis["warnings"]))

    # 6. Resume with no Work Experience section
    def test_06_resume_with_no_work_experience(self):
        text = """Alice Student
        Education:
        B.S. in Computer Science, 2024
        
        Skills:
        Python, Java, React, SQL
        
        Projects:
        Mini Search Engine
        • Built TF-IDF crawler in Python
        """
        analysis = ai_service.analyze_resume(text, "Student.txt", file_type="txt")
        profile = analysis["extracted_profile"]
        self.assertEqual(len(profile["experience"]), 0)
        self.assertEqual(len(profile["projects"]), 1)
        self.assertIn("Python", profile["skills"])

    # 7. PDF containing selectable text
    def test_07_pdf_selectable_text(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 50), "Dev Candidate\nProfessional Summary\nExperienced engineer.\nSkills\nPython, Docker\nProjects\nMicroservice Gateway\n• Built FastAPI gateway with Redis")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        text, ftype = extract_text_from_file("selectable.pdf", pdf_bytes)
        self.assertEqual(ftype, "pdf")
        self.assertIn("Microservice Gateway", text)

        analysis = ai_service.analyze_resume(text, "selectable.pdf", file_type="pdf")
        self.assertEqual(len(analysis["extracted_profile"]["projects"]), 1)

    # 8. Image-only/scanned PDF (detects and reports OCR required)
    def test_08_image_only_scanned_pdf(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        # Insert bitmap without selectable text
        pix = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 80, 80), 0)
        page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pix)
        pdf_bytes = pdf.tobytes()
        pdf.close()

        with self.assertRaises(ResumeParserError) as ctx:
            extract_text_from_file("scanned.pdf", pdf_bytes)
        self.assertIn("OCR is required", str(ctx.exception))

    # 9. Empty, malformed, unsupported, or corrupted uploads
    def test_09_empty_malformed_unsupported(self):
        # Empty bytes
        with self.assertRaises(ResumeParserError):
            extract_text_from_file("empty.docx", b"")

        # Unsupported extension
        with self.assertRaises(ResumeParserError):
            extract_text_from_file("virus.exe", b"binary content")

        # Corrupted DOCX (not a zip)
        with self.assertRaises(ResumeParserError):
            extract_text_from_file("corrupt.docx", b"not a real docx file")

        # Corrupted PDF
        with self.assertRaises(ResumeParserError):
            extract_text_from_file("corrupt.pdf", b"garbage pdf data")

    # 10. Multiple-page resume with section content crossing page boundaries
    def test_10_multiple_page_resume(self):
        pdf = pymupdf.open()
        p1 = pdf.new_page()
        p1.insert_text((50, 50), "Candidate MultiPage\nTechnical Skills\nPython, PostgreSQL, React\nProjects\nInventory Management System\n• Real-time stock alerts")
        p2 = pdf.new_page()
        p2.insert_text((50, 50), "• Implemented barcode scanning with TypeScript\n• Automated reordering pipelines\nEducation\nB.S. Software Engineering")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        text, _ = extract_text_from_file("multipage.pdf", pdf_bytes)
        analysis = ai_service.analyze_resume(text)
        projects = analysis["extracted_profile"]["projects"]
        self.assertGreaterEqual(len(projects), 1)
        inv_proj = projects[0]
        # Verify evidence across pages was captured
        self.assertGreaterEqual(len(inv_proj["evidence"]), 2)

    # 11. Resume containing unknown sections that must not be discarded
    def test_11_unknown_sections_not_discarded(self):
        text = """Dr. Researcher
        Professional Summary:
        AI Research Scientist.
        
        Skills:
        PyTorch, TensorFlow, Python
        
        Publications:
        - Deep Residual Learning for Signal Detection, IEEE 2023
        - Transformer Attention Scaling, NeurIPS 2024
        
        Awards & Honors:
        - Best Paper Award, CVPR 2022
        """
        detailed = extract_sections_detailed(text)
        unrec = detailed["unrecognized"]
        self.assertTrue(any("Publications" in k for k in unrec), f"Publications missing from {unrec.keys()}")
        self.assertTrue(any("Awards" in k for k in unrec), f"Awards missing from {unrec.keys()}")

    # 12. Resume with similar skill names appearing in unrelated sections
    def test_12_similar_skills_in_unrelated_sections(self):
        text = """Engineer Test
        Professional Summary:
        Passionate Python developer interested in cloud infrastructure.
        
        Technical Skills:
        Python, Docker, SQL
        
        Projects:
        Analytics Bot
        • Developed data pipeline using SQL
        """
        analysis = ai_service.analyze_resume(text)
        projects = analysis["extracted_profile"]["projects"]
        self.assertEqual(len(projects), 1)
        # Bot project should have SQL, but not claim Docker just because it was in skills
        bot_techs = projects[0]["technologies"]
        self.assertIn("SQL", bot_techs)
        self.assertNotIn("Docker", bot_techs)

    # 13. Partial extraction and parser exception scenarios
    def test_13_partial_extraction_detection(self):
        # Text has a prominent PROJECTS heading but 0 parseable project bullets
        text = """Candidate X
        Professional Summary:
        Developer
        
        Projects:
        ???
        
        Education:
        B.S. in Computer Science
        """
        analysis = ai_service.analyze_resume(text)
        # Should identify that projects section existed in source but nothing was structured
        self.assertEqual(analysis["extraction_status"], "partial")
        self.assertTrue(any("Project section detected" in w for w in analysis["warnings"]))

    # 14. Ensure AI-generated text cannot be mistaken for source evidence
    def test_14_ai_text_not_mistaken_for_source_evidence(self):
        text = "Candidate Y\nSkills:\nPython, React\nProjects:\nChat App\n• Built real-time socket server in Python"
        analysis = ai_service.analyze_resume(text)
        p = analysis["extracted_profile"]["projects"][0]

        # Inferred skills must be cleanly segregated from explicit technologies
        self.assertIn("Python", p["technologies"])
        self.assertNotIn("FastAPI", p["technologies"])
        # Source text must equal input
        self.assertIn("Built real-time socket server", analysis["raw_text"])

    # 15. Verify downstream matching receives preserved project evidence
    def test_15_matching_receives_preserved_project_evidence(self):
        docx_bytes = build_test_cv_docx_bytes()
        text, _ = extract_text_from_file("CareerLens_Test_CV.docx", docx_bytes)
        analysis = ai_service.analyze_resume(text)
        profile = analysis["extracted_profile"]

        job = {
            "id": "job-fastapi-backend",
            "title": "FastAPI Backend Developer",
            "requirements": [
                {"name": "FastAPI", "importance": "critical"},
                {"name": "PostgreSQL", "importance": "critical"},
                {"name": "REST APIs", "importance": "critical"}
            ],
            "description": "High throughput backend APIs using Python, FastAPI, and PostgreSQL."
        }

        match = matching_service.match_resume_to_job(profile, text, job)
        self.assertGreaterEqual(match["project_relevance"], 70)
        
        # Verify evidence matched from projects
        matched = {m["requirement"]: m for m in match["matched_requirements"]}
        self.assertIn("FastAPI", matched)
        self.assertEqual(matched["FastAPI"]["match_type"], "project_evidence")
        self.assertIn("TaskFlow", matched["FastAPI"]["resume_evidence"])

    # 16. Verify existing upload, analysis, and matching flows continue working
    def test_16_existing_api_upload_and_matching_flow(self):
        docx_bytes = build_test_cv_docx_bytes()
        
        # Test /api/resume/upload endpoint
        files = {"file": ("CareerLens_Test_CV.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = self.client.post("/api/resume/upload", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["extraction_status"], "warning")  # Has 'Additional Information' preserved
        self.assertEqual(len(data["extracted_profile"]["projects"]), 3)

        # Test /api/match with uploaded resume
        match_res = self.client.post("/api/match", json={
            "resume_id": data["resume_id"],
            "job_id": "job-backend-001"
        })
        self.assertEqual(match_res.status_code, 200)
        match_data = match_res.json()
        self.assertGreaterEqual(match_data["overall_score"], 50)
        self.assertGreaterEqual(match_data["project_relevance"], 70)


if __name__ == "__main__":
    unittest.main()
