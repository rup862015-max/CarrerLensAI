"""
Resume Parser Service for CareerLens AI
Extracts plain text, preserves logical structure, headings, and bullet points from PDF and DOCX files.
Handles error cases like corrupt files, empty content, image-only/scanned documents, and unsupported formats.
"""

import io
import re
from typing import Dict, List, Any, Tuple, Optional
import pymupdf  # PyMuPDF
import docx


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB limit


class ResumeParserError(Exception):
    """Custom exception for resume parsing errors."""
    pass


# Canonical section headings patterns (case-insensitive, matched on stripped line)
CANONICAL_PATTERNS = {
    "summary": r"^(?:professional\s+summary|executive\s+summary|summary|profile|professional\s+profile|career\s+profile|about\s+me|career\s+objective|objective|personal\s+statement)$",
    "skills": r"^(?:technical\s+skills|core\s+skills|skills\s+(?:&|and)\s+technologies|skills\s+(?:&|and)\s+tools|skills|technologies|core\s+competencies|competencies|technical\s+proficiencies|technical\s+competencies|areas\s+of\s+expertise|expertise)$",
    "projects": r"^(?:projects|technical\s+projects|academic\s+projects|personal\s+projects|selected\s+projects|key\s+projects|recent\s+projects|engineering\s+projects|software\s+projects|portfolio)$",
    "experience": r"^(?:work\s+experience|professional\s+experience|employment\s+history|experience|career\s+history|work\s+history|professional\s+background|relevant\s+experience|internships|industry\s+experience)$",
    "education": r"^(?:education|academic\s+background|academics|academic\s+history|educational\s+background|qualifications|degrees|educational\s+qualifications)$",
    "certifications": r"^(?:certifications|certificates|licenses|professional\s+certifications|certifications\s+(?:&|and)\s+licenses|courses)$"
}

# Sub-headings inside a project or experience entry that must NEVER be treated as top-level section headings
SUB_HEADING_PATTERN = r"^(?:technologies|technology\s+details|tech\s+stack|tools\s+used|tools|environment|key\s+highlights|role|deliverables|responsibilities|achievements\s+summary|duration|client|github|link|repository|metrics)$"

# Known common non-canonical / additional section headings
ADDITIONAL_SECTION_PATTERNS = r"^(?:additional\s+information|additional\s+details|other\s+information|publications|patents|volunteer\s+experience|volunteering|volunteer\s+work|languages|affiliations|extracurricular\s+activities|activities|interests|hobbies|references|awards|honors|awards\s+(?:&|and)\s+honors|achievements)$"


def clean_text(text: str) -> str:
    """
    Normalizes whitespace and standardizes bullet characters without corrupting punctuation.
    Preserves em-dashes, en-dashes, and quotes.
    Standardizes unicode bullets and PyMuPDF middle-dots into '• '.
    """
    if not text:
        return ""

    # Normalize line breaks
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    
    # Strip non-printable ASCII control characters (0x00-0x08, 0x0B, 0x0C, 0x0E-0x1F)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', text)
    
    # Standardize common unicode bullets, middle dots, dingbats, and replacement chars into '• '
    # Including \xb7 (middle dot), \u2022, \u2023, \u25e6, \u2043, \u2219, \uf0b7, \uf0a7, \ufffd, \x95
    text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219•●■◆▪\uf0b7\uf0a7\xb7·\ufffd\x95]\s*', '• ', text)

    # Clean whitespace line by line
    cleaned_lines = []
    for line in text.split('\n'):
        # Normalize internal whitespace
        line_clean = re.sub(r'[ \t]+', ' ', line).strip()
        cleaned_lines.append(line_clean)

    # Rejoin lines
    result = '\n'.join(cleaned_lines)
    
    # Collapse 3+ consecutive newlines to double newlines
    result = re.sub(r'\n{3,}', '\n\n', result)
    return result.strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts raw text from a PDF byte buffer using PyMuPDF with layout and reading order preservation."""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        if doc.page_count == 0:
            raise ResumeParserError("The uploaded PDF has 0 pages.")
        
        full_text = []
        total_images = 0
        for page_num in range(doc.page_count):
            page = doc.load_page(page_num)
            # Use block extraction to preserve reading order
            blocks = page.get_text("blocks")
            # Sort blocks by vertical position then horizontal
            blocks.sort(key=lambda b: (b[1], b[0]))
            page_text = "\n".join(b[4].strip() for b in blocks if len(b) > 4 and b[4].strip())
            if page_text:
                full_text.append(page_text)
            total_images += len(page.get_images())
        
        doc.close()
        combined = "\n\n".join(full_text)
        cleaned = clean_text(combined)

        # Scanned / Image-only document detection
        if not cleaned:
            if total_images > 0:
                raise ResumeParserError("This document appears to contain limited readable text. It may be scanned or image-based. OCR is required.")
            raise ResumeParserError("The uploaded PDF contains no readable text.")

        if len(cleaned.split()) < 10 and total_images > 0:
            raise ResumeParserError("This document appears to contain limited readable text. It may be scanned or image-based. OCR is required.")

        return cleaned
    except Exception as e:
        if isinstance(e, ResumeParserError):
            raise e
        raise ResumeParserError(f"Failed to parse PDF file: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from a DOCX byte buffer using python-docx.
    Preserves paragraph reading order, section headings, list bullet markers, and tables.
    """
    try:
        file_stream = io.BytesIO(file_bytes)
        doc = docx.Document(file_stream)
        
        full_text = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Detect if this paragraph is a list / bullet style in Word
            is_bullet_style = False
            if para.style and para.style.name:
                style_name = para.style.name.lower()
                if "bullet" in style_name or "list" in style_name:
                    is_bullet_style = True

            # Also check underlying Word XML numbering definition (w:numPr)
            if not is_bullet_style and hasattr(para, "_element") and para._element.pPr is not None:
                is_bullet_style = bool(para._element.pPr.xpath("./w:numPr"))

            # Prepend bullet if it's a list item and doesn't already have a bullet character
            if is_bullet_style and not text.startswith(('•', '-', '*', '1.', '2.', '3.')):
                text = f"• {text}"

            full_text.append(text)
                
        # Extract tables while preserving row structure
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        combined = "\n".join(full_text)
        if not combined.strip():
            raise ResumeParserError("The uploaded DOCX file is empty.")
        return clean_text(combined)
    except Exception as e:
        if isinstance(e, ResumeParserError):
            raise e
        raise ResumeParserError(f"Failed to parse DOCX file: {str(e)}")


def extract_text_from_file(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Main parser entry point.
    Returns (extracted_text, file_type).
    """
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ResumeParserError("File exceeds maximum allowed size of 10 MB.")
        
    if not file_bytes:
        raise ResumeParserError("Uploaded file is empty (0 bytes).")
        
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes), "pdf"
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes), "docx"
    elif lower_name.endswith(".txt"):
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            cleaned = clean_text(text)
            if not cleaned.strip():
                raise ResumeParserError("The uploaded text file is empty.")
            return cleaned, "txt"
        except Exception as e:
            if isinstance(e, ResumeParserError):
                raise e
            raise ResumeParserError(f"Failed to read TXT file: {str(e)}")
    else:
        raise ResumeParserError(f"Unsupported file format '{filename}'. Please upload a PDF (.pdf) or Word (.docx) document.")


def extract_sections_detailed(text: str) -> Dict[str, Any]:
    """
    Splits resume text into canonical and unrecognized sections with rigorous line-by-line heading detection.
    Returns:
      {
        "canonical": { "summary": str, "experience": str, "skills": str, "education": str, "projects": str, "certifications": str, "other": str },
        "unrecognized": { "Section Name": str, ... },
        "header": str (content preceding the first heading),
        "section_order": List[str]
      }
    """
    canonical_sections = {
        "summary": "",
        "experience": "",
        "skills": "",
        "education": "",
        "projects": "",
        "certifications": "",
        "other": ""
    }
    unrecognized_sections: Dict[str, str] = {}
    section_order: List[str] = []

    if not text or not text.strip():
        return {
            "canonical": canonical_sections,
            "unrecognized": unrecognized_sections,
            "header": "",
            "section_order": section_order
        }

    lines = text.split('\n')
    
    # Identify headings line by line
    heading_matches = []  # List of (line_idx, section_key, display_title, is_unrecognized)
    
    for idx, raw_line in enumerate(lines):
        line = raw_line.strip()
        if not line:
            continue

        # Headings cannot start with bullet points or list enumeration
        if line.startswith(('•', '-', '*', '1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.')):
            continue

        candidate = line.rstrip(':').strip()
        # Headings are concise titles, not sentences
        words = candidate.split()
        if len(candidate) > 55 or len(words) > 6 or len(candidate) < 2:
            continue

        c_lower = candidate.lower()

        # Reject sub-headings like 'Technologies:', 'Tech Stack:', etc.
        if re.match(SUB_HEADING_PATTERN, c_lower):
            continue

        # Check for canonical section match
        matched_canonical = None
        for sec_name, pattern in CANONICAL_PATTERNS.items():
            if re.match(pattern, c_lower):
                matched_canonical = sec_name
                break

        if matched_canonical:
            heading_matches.append((idx, matched_canonical, candidate, False))
        elif re.match(ADDITIONAL_SECTION_PATTERNS, c_lower):
            heading_matches.append((idx, candidate.title(), candidate, True))
        elif candidate.isupper() and len(words) <= 4 and not any(char.isdigit() for char in candidate):
            # All-caps line that looks like a custom heading (e.g., "AWARDS & HONORS")
            heading_matches.append((idx, candidate.title(), candidate, True))

    if not heading_matches:
        # Fallback: estimate first block as summary, remaining as experience
        canonical_sections["summary"] = text[:300].strip()
        canonical_sections["experience"] = text[300:].strip()
        return {
            "canonical": canonical_sections,
            "unrecognized": unrecognized_sections,
            "header": "",
            "section_order": ["summary", "experience"]
        }

    # Slice header text before the first section
    first_heading_idx = heading_matches[0][0]
    header_text = "\n".join(lines[:first_heading_idx]).strip()

    # Slice each section
    for i, (line_idx, sec_key, display_title, is_unrec) in enumerate(heading_matches):
        next_line_idx = heading_matches[i + 1][0] if i + 1 < len(heading_matches) else len(lines)
        # Content starts from line_idx + 1 up to next_line_idx
        content = "\n".join(lines[line_idx + 1:next_line_idx]).strip()

        if is_unrec:
            if sec_key in unrecognized_sections:
                unrecognized_sections[sec_key] = f"{unrecognized_sections[sec_key]}\n\n{content}".strip()
            else:
                unrecognized_sections[sec_key] = content
                section_order.append(sec_key)
        else:
            if canonical_sections[sec_key]:
                # Combine multiple occurrences rather than overwriting!
                canonical_sections[sec_key] = f"{canonical_sections[sec_key]}\n\n{content}".strip()
            else:
                canonical_sections[sec_key] = content
                section_order.append(sec_key)

    return {
        "canonical": canonical_sections,
        "unrecognized": unrecognized_sections,
        "header": header_text,
        "section_order": section_order
    }


def extract_sections(text: str) -> Dict[str, str]:
    """
    Backward-compatible wrapper for extract_sections.
    Returns standard dictionary with canonical keys and an 'unrecognized_sections' mapping.
    """
    detailed = extract_sections_detailed(text)
    sections = detailed["canonical"].copy()
    sections["unrecognized_sections"] = detailed["unrecognized"]
    return sections
