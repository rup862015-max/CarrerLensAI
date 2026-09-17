"""
Comprehensive Unit & Integration Test Suite for CareerLens AI Phase 1
Tests parsing, skill extraction, matching engine, explainability, SQLite persistence, and API endpoints.
"""

import os
import sys
import unittest
import json
import pymupdf
from fastapi.testclient import TestClient

# Add current directory to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from database import init_db, get_db_connection, save_resume, save_resume_analysis, get_latest_resume
from services.resume_parser import clean_text, extract_sections, extract_text_from_file
from services.skill_service import skill_service
from services.explanation_service import explanation_service
from services.matching_service import matching_service
from services.ai_service import ai_service
from main import app, DEMO_RESUME_TEXT


class TestCareerLensPhase1(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        init_db()
        cls.client = TestClient(app)

    def test_01_clean_text_and_sections(self):
        sample = "John Doe\n\nPROFESSIONAL SUMMARY\nExperienced developer\n\nTECHNICAL SKILLS\nPython, SQL, React"
        cleaned = clean_text(sample)
        self.assertIn("John Doe", cleaned)
        sections = extract_sections(cleaned)
        self.assertTrue("skills" in sections or "summary" in sections)

    def test_02_skill_service(self):
        text = "Experienced in Python, React, PostgreSQL, Docker, and AWS."
        skills = skill_service.extract_skills(text)
        names = [s["name"] for s in skills]
        self.assertIn("Python", names)
        self.assertIn("React", names)
        self.assertIn("PostgreSQL", names)
        self.assertIn("Docker", names)
        self.assertIn("AWS", names)

        # Related skills
        related = skill_service.get_related_skills("React")
        self.assertTrue(len(related) > 0)

    def test_03_ai_service_resume_analysis(self):
        analysis = ai_service.analyze_resume(DEMO_RESUME_TEXT, "Alex_Mercer_Resume.pdf")
        self.assertEqual(analysis["filename"], "Alex_Mercer_Resume.pdf")
        
        health = analysis["health_scores"]
        self.assertGreaterEqual(health["overall"], 70)
        self.assertGreaterEqual(health["skills_coverage"], 70)

        profile = analysis["extracted_profile"]
        self.assertIn("Python", profile["skills"])
        self.assertIn("React", profile["skills"])
        self.assertTrue(len(profile["projects"]) >= 2)

    def test_04_matching_engine_backend_vs_data_analyst(self):
        # Analyze demo resume
        analysis = ai_service.analyze_resume(DEMO_RESUME_TEXT, "Alex_Mercer_Resume.pdf")
        profile = analysis["extracted_profile"]

        # Load demo jobs
        res = self.client.get("/api/jobs/demo")
        self.assertEqual(res.status_code, 200)
        jobs = res.json()
        self.assertEqual(len(jobs), 3)

        backend_job = next(j for j in jobs if "Backend" in j["title"])
        data_job = next(j for j in jobs if "Data" in j["title"])

        # Match Backend Developer
        backend_match = matching_service.match_resume_to_job(profile, DEMO_RESUME_TEXT, backend_job)
        # Match Data Analyst
        data_match = matching_service.match_resume_to_job(profile, DEMO_RESUME_TEXT, data_job)

        # Backend match score should be higher than Data Analyst for this candidate
        self.assertGreater(backend_match["overall_score"], data_match["overall_score"])
        self.assertGreaterEqual(backend_match["overall_score"], 80)

        # Explainability check
        matched_req_names = [m["requirement"] for m in backend_match["matched_requirements"]]
        self.assertIn("Python", matched_req_names)
        self.assertIn("REST APIs", matched_req_names)

        # Verify evidence is present
        for m in backend_match["matched_requirements"]:
            self.assertTrue(len(m["resume_evidence"]) > 0)
            self.assertGreater(m["confidence"], 0.5)

    def test_05_api_health(self):
        res = self.client.get("/api/health")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["status"], "healthy")

    def test_06_api_resume_analyze_and_match_flow(self):
        # 1. Post analyze text
        res = self.client.post("/api/resume/analyze", json={
            "resume_text": DEMO_RESUME_TEXT,
            "filename": "Alex_Mercer.txt"
        })
        self.assertEqual(res.status_code, 200)
        resume_id = res.json()["resume_id"]

        # 2. Get latest resume
        res_latest = self.client.get("/api/resume/latest")
        self.assertEqual(res_latest.status_code, 200)
        self.assertEqual(res_latest.json()["resume_id"], resume_id)

        # 3. Post match with demo job id
        match_res = self.client.post("/api/match", json={
            "resume_id": resume_id,
            "job_id": "job-backend-001"
        })
        self.assertEqual(match_res.status_code, 200)
        match_data = match_res.json()
        self.assertIn("overall_score", match_data)
        self.assertIn("matched_requirements", match_data)
        self.assertIn("skill_gaps", match_data)
        self.assertIn("career_preview", match_data)

    def test_07_pdf_and_docx_parsing(self):
        import docx
        import io

        # 1. Test DOCX generation & parsing
        doc = docx.Document()
        doc.add_heading("Alex Mercer", level=0)
        doc.add_paragraph("PROFESSIONAL SUMMARY\nSoftware Engineer with experience in Python and React.")
        doc.add_paragraph("TECHNICAL SKILLS\nPython, React, Node.js, SQL, Docker")
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        docx_bytes = doc_stream.getvalue()

        text, ftype = extract_text_from_file("test_resume.docx", docx_bytes)
        self.assertEqual(ftype, "docx")
        self.assertIn("Alex Mercer", text)
        self.assertIn("Python", text)

        # 2. Test PDF generation & parsing
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page()
        page.insert_text((50, 50), "Alex Mercer\nSoftware Engineer\nSkills: Python, TypeScript, FastAPI")
        pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()

        ptext, pftype = extract_text_from_file("test_resume.pdf", pdf_bytes)
        self.assertEqual(pftype, "pdf")
        self.assertIn("Alex Mercer", ptext)
        self.assertIn("FastAPI", ptext)


if __name__ == "__main__":
    unittest.main()
